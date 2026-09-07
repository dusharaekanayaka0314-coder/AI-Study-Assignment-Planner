import os
import docx
from dotenv import load_dotenv
import google.generativeai as genai
import PyPDF2
import streamlit as st

# =========================================================
# CONFIG & INITIALIZATION
# =========================================================
load_dotenv()
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

genai.configure(api_key=GEMINI_API_KEY)
model = genai.GenerativeModel("gemini-3.6-flash")

st.set_page_config(
    page_title="AI Study & Assignment Planner",
    page_icon="🎓",
    layout="centered"
)

# =========================================================
# MODERN UI STYLING
# =========================================================
st.markdown(
    """
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');

/* Apply modern font across application */
html, body, [class*="css"], p, span, label, div, h1, h2, h3, h4, h5, h6 {
    font-family: 'Inter', sans-serif !important;
}

/* Background & Body Colors */
.stApp {
    background-color: #FAFAFC !important;
}

/* Field Label Styling */
label, .stMarkdown label p {
    font-size: 0.92rem !important;
    font-weight: 600 !important;
    color: #2D3748 !important;
}

/* Main Header Card */
.main-header {
    text-align: center;
    padding: 36px 24px;
    background: linear-gradient(135deg, #6366F1 0%, #4F46E5 100%);
    border-radius: 20px;
    margin-bottom: 28px;
    box-shadow: 0 10px 25px -5px rgba(99, 102, 241, 0.25);
    color: #FFFFFF;
}

.main-header h1 {
    color: #FFFFFF !important;
    font-size: 2.1rem !important;
    font-weight: 700 !important;
    margin-bottom: 10px !important;
    letter-spacing: -0.02em;
}

.main-header p {
    color: #E0E7FF !important;
    font-size: 1.05rem !important;
    font-weight: 400 !important;
    max-width: 580px;
    margin: 0 auto;
}

/* Modern Card Headers */
.card-title {
    display: flex;
    align-items: center;
    gap: 10px;
    font-size: 1.2rem;
    font-weight: 700;
    color: #1E293B;
    margin-bottom: 18px;
}

.card-title svg {
    color: #6366F1;
}

/* Primary Action Button */
div.stButton > button {
    background: linear-gradient(135deg, #6366F1 0%, #4F46E5 100%) !important;
    color: #FFFFFF !important;
    font-weight: 600 !important;
    font-size: 1.05rem !important;
    border-radius: 12px !important;
    padding: 14px 28px !important;
    border: none !important;
    width: 100% !important;
    box-shadow: 0 4px 14px rgba(79, 70, 229, 0.3) !important;
    transition: all 0.2s ease-in-out !important;
    margin-top: 10px;
}

div.stButton > button:hover {
    transform: translateY(-1px) !important;
    box-shadow: 0 6px 20px rgba(79, 70, 229, 0.4) !important;
}

/* Output Container Styling */
.output-box {
    background-color: #FFFFFF;
    border: 1px solid #E2E8F0;
    border-radius: 16px;
    padding: 28px;
    box-shadow: 0 4px 12px rgba(0, 0, 0, 0.03);
    margin-top: 24px;
}

/* Expanders */
.streamlit-expanderHeader {
    font-weight: 600 !important;
    color: #475569 !important;
}
</style>
""",
    unsafe_allow_html=True,
)

# =========================================================
# HEADER
# =========================================================
st.markdown(
    """
<div class="main-header">
    <div style="font-size: 3rem; margin-bottom: 8px;">🎓</div>
    <h1>AI Study & Assignment Planner</h1>
    <p>Upload your syllabus or assignment guidelines to generate an actionable, day-by-day study roadmap.</p>
</div>
""",
    unsafe_allow_html=True,
)

if "extracted_text" not in st.session_state:
    st.session_state.extracted_text = ""

# =========================================================
# BASIC DETAILS SECTION
# =========================================================
with st.container():
    st.markdown(
        """
    <div class="card-title">
        <svg width="22" height="22" fill="none" stroke="currentColor" stroke-width="2" viewBox="0 0 24 24">
            <path stroke-linecap="round" stroke-linejoin="round" d="M16 7a4 4 0 11-8 0 4 4 0 018 0zM12 14a7 7 0 00-7 7h14a7 7 0 00-7-7z"></path>
        </svg>
        Basic Details
    </div>
    """,
        unsafe_allow_html=True,
    )

    col1, col2 = st.columns(2)

    with col1:
        subject = st.text_input("Subject Name", placeholder="e.g. Data Structures")
        study_hours = st.number_input(
            "Study Hours per Day", min_value=1, max_value=16, value=2
        )

    with col2:
        deadline = st.date_input("Exam / Assignment Deadline")
        level = st.selectbox(
            "Current Level", ["Beginner", "Intermediate", "Advanced"]
        )

    goal = st.text_area(
        "Goal",
        placeholder="e.g., Score above 85%, Master core concepts, Finish programming task",
        height=100
    )

st.write("")

# =========================================================
# SYLLABUS / ASSIGNMENT MATERIAL SECTION
# =========================================================
with st.container():
    st.markdown(
        """
    <div class="card-title">
        <svg width="22" height="22" fill="none" stroke="currentColor" stroke-width="2" viewBox="0 0 24 24">
            <path stroke-linecap="round" stroke-linejoin="round" d="M9 12h6m-6 4h6m2 5H7a2 2 0 01-2-2V5a2 2 0 012-2h5.586a1 1 0 01.707.293l5.414 5.414a1 1 0 01.293.707V19a2 2 0 01-2 2z"></path>
        </svg>
        Syllabus / Assignment Material
    </div>
    """,
        unsafe_allow_html=True,
    )

    input_method = st.radio(
        "Select Input Method:",
        ["Upload File (PDF/DOCX/TXT)", "Paste Text"],
        horizontal=True
    )

    if input_method == "Upload File (PDF/DOCX/TXT)":
        uploaded_file = st.file_uploader(
            "Upload your study document", type=["pdf", "docx", "txt"]
        )

        if uploaded_file is not None:
            text = ""

            if uploaded_file.name.endswith(".pdf"):
                reader = PyPDF2.PdfReader(uploaded_file)
                for page in reader.pages:
                    text += page.extract_text() or ""

            elif uploaded_file.name.endswith(".docx"):
                doc = docx.Document(uploaded_file)
                for para in doc.paragraphs:
                    text += para.text + "\n"
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            text += cell.text + " "
                        text += "\n"

            elif uploaded_file.name.endswith(".txt"):
                text = uploaded_file.read().decode("utf-8")

            st.session_state.extracted_text = text

            if text.strip():
                st.success("File processed successfully.")
                with st.expander("Preview Extracted Text"):
                    st.write(
                        text[:2000] + ("..." if len(text) > 2000 else "")
                    )
            else:
                st.warning(
                    "No text could be extracted from this file. Try another file."
                )
    else:
        pasted = st.text_area(
            "Paste your syllabus or instructions here", height=180
        )
        st.session_state.extracted_text = pasted

st.write("")

# =========================================================
# GENERATE PLAN
# =========================================================
if st.button("Generate Study Plan"):
    if not st.session_state.extracted_text.strip():
        st.error("Please upload a file or paste text first.")
    else:
        with st.spinner("Analyzing material and generating your personalized schedule..."):
            prompt = f"""
You are a study planner AI.

Based on the details and material below, respond using EXACTLY this format with markdown headers:

## Important Topics
(bullet list)

## High-Priority Topics
(bullet list)

## Study Plan / Timetable
(day-by-day breakdown as a bullet list or table)

## What to Focus On
(2-3 sentences)

Subject: {subject}
Deadline: {deadline}
Study Hours per Day: {study_hours}
Current Level: {level}
Goal: {goal}

Material:
{st.session_state.extracted_text[:5000]}
"""
            response = model.generate_content(prompt)

        st.markdown('<div class="output-box">', unsafe_allow_html=True)
        st.markdown(response.text)
        st.markdown('</div>', unsafe_allow_html=True)
